from __future__ import annotations

import time
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import LedgerRecord

HEADERS = [
    "Internal ID",
    "Source Type",
    "Source Email Message ID",
    "Sender Email",
    "Original File Name",
    "Stored File Name",
    "Drive File ID",
    "Drive Folder Path",
    "Vendor Name",
    "Invoice Number",
    "Invoice Date",
    "Due Date",
    "Currency",
    "Subtotal",
    "Tax Amount",
    "Total Amount",
    "Category",
    "Extraction Confidence",
    "Status",
    "Review Notes",
    "Created At",
    "Updated At",
]


def ensure_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(HEADERS)
    wb.save(path)


def append_or_update_record(path: Path, record: LedgerRecord, retries: int = 3) -> None:
    ensure_workbook(path)
    for attempt in range(retries):
        try:
            wb = load_workbook(path)
            ws = wb["Invoices"]
            id_col = 1
            row_to_update = None
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=id_col).value == record.internal_id:
                    row_to_update = row
                    break
            values = [
                record.internal_id,
                record.source_type,
                record.source_email_message_id,
                record.sender_email,
                record.original_file_name,
                record.stored_file_name,
                record.drive_file_id,
                record.drive_folder_path,
                record.vendor_name,
                record.invoice_number,
                record.invoice_date,
                record.due_date,
                record.currency,
                record.subtotal,
                record.tax_amount,
                record.total_amount,
                record.category,
                record.extraction_confidence,
                record.status,
                record.review_notes,
                record.created_at,
                record.updated_at,
            ]
            if row_to_update:
                for col, value in enumerate(values, start=1):
                    ws.cell(row=row_to_update, column=col, value=value)
            else:
                ws.append(values)
            wb.save(path)
            return
        except PermissionError:
            if attempt == retries - 1:
                raise
            time.sleep(1.0)


def reconcile(path: Path) -> tuple[int, int]:
    ensure_workbook(path)
    wb = load_workbook(path)
    ws = wb["Invoices"]
    seen = set()
    dup = 0
    for row in range(2, ws.max_row + 1):
        internal_id = ws.cell(row=row, column=1).value
        if not internal_id:
            continue
        if internal_id in seen:
            dup += 1
        seen.add(internal_id)
    return ws.max_row - 1, dup
